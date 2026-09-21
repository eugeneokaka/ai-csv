import helper
from openpyxl.styles import PatternFill

df = helper.get_full_csv("data_without_city.csv")

helper.save_excel(df, "salary_green.xlsx")

wb = helper.load_workbook("salary_green.xlsx")
ws = wb.active

green_fill = PatternFill(start_color="92D050", end_color="92D050", fill_type="solid")

for row in range(2, ws.max_row + 1):
    ws.cell(row=row, column=6).fill = green_fill

wb.save(helper.get_output_path("salary_green.xlsx"))

print("Salary column highlighted in green!")
print(df.to_string(index=False))